"""Kiểm thử ô bảng gộp cột (colspan) và gộp hàng (rowspan) cùng việc ẩn đường kẻ viền bên trong."""
import pytest

from chuviettay.document.ir import Paragraph, Table, TableBorder, TableCell, TableRow, Text
from chuviettay.layout.table_layout import TableLayoutEngine


def test_table_colspan_measurement_and_border_suppression():
    # Bảng 2 hàng, 2 cột: Hàng 1 gộp 2 cột (colspan=2), Hàng 2 gồm 2 ô riêng biệt
    row1 = TableRow(cells=[TableCell(blocks=[Paragraph(inlines=[Text("Tiêu đề gộp")])], colspan=2)])
    row2 = TableRow(cells=[TableCell.from_text("Ô 1"), TableCell.from_text("Ô 2")])
    table = Table(rows=[row1, row2], border_style=TableBorder.ALL)

    engine = TableLayoutEngine(available_width=400.0, line_height=20.0)
    data = engine.layout_table(table, x0=50.0, y0=50.0)

    assert len(data.col_widths) == 2
    # Ô gộp hàng 1 phải có bề rộng bằng tổng 2 cột
    assert abs(data.cells[0][0].width - sum(data.col_widths)) < 0.1

    # Nét viền sinh ra không được có đường phân cách dọc cắt ngang qua hàng 1
    strokes = engine.generate_border_strokes(data, TableBorder.ALL)
    mid_x = data.x + data.col_widths[0]
    row1_top_y = data.y
    row1_bottom_y = data.y + data.row_heights[0]

    # Kiểm tra không có nét dọc nào đi qua khoảng từ row1_top_y đến row1_bottom_y tại mid_x
    for s in strokes:
        pts = s.points
        if len(pts) == 2 and abs(pts[0][0] - mid_x) < 0.5 and abs(pts[1][0] - mid_x) < 0.5:
            # Đường dọc tại vị trí giữa 2 cột: không được nằm trọn trong hàng 1
            min_y = min(pts[0][1], pts[1][1])
            max_y = max(pts[0][1], pts[1][1])
            assert not (min_y >= row1_top_y - 0.1 and max_y <= row1_bottom_y + 0.1), (
                f"Đường kẻ dọc cắt qua ô gộp colspan tại x={mid_x}"
            )


def test_table_rowspan_measurement_and_border_suppression():
    # Bảng 2 hàng, 2 cột: Cột 1 gộp 2 hàng (rowspan=2)
    row1 = TableRow(cells=[
        TableCell(blocks=[Paragraph(inlines=[Text("Gộp hàng")])], rowspan=2),
        TableCell.from_text("H1 C2"),
    ])
    row2 = TableRow(cells=[
        TableCell.from_text("H2 C2"),
    ])
    table = Table(rows=[row1, row2], border_style=TableBorder.ALL)

    engine = TableLayoutEngine(available_width=400.0, line_height=20.0)
    data = engine.layout_table(table, x0=50.0, y0=50.0)

    # Ô gộp cột 1 phải có chiều cao bằng tổng 2 hàng
    assert abs(data.cells[0][0].height - sum(data.row_heights)) < 0.1

    # Nét viền sinh ra không được có đường ngang cắt qua cột 1 giữa hàng 1 và hàng 2
    strokes = engine.generate_border_strokes(data, TableBorder.ALL)
    divider_y = data.y + data.row_heights[0]
    col0_left_x = data.x
    col0_right_x = data.x + data.col_widths[0]

    for s in strokes:
        pts = s.points
        if len(pts) == 2 and abs(pts[0][1] - divider_y) < 0.5 and abs(pts[1][1] - divider_y) < 0.5:
            min_x = min(pts[0][0], pts[1][0])
            max_x = max(pts[0][0], pts[1][0])
            assert not (min_x >= col0_left_x - 0.1 and max_x <= col0_right_x + 0.1), (
                f"Đường kẻ ngang cắt qua ô gộp rowspan tại y={divider_y}"
            )


def test_docx_import_table_with_merged_cells(tmp_path):
    import docx
    from chuviettay.importer.docx_importer import DocxImporter

    doc_path = str(tmp_path / "merged_table.docx")
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Tiêu đề gộp"
    t.cell(0, 0).merge(t.cell(0, 1))
    t.cell(1, 0).text = "Ô 1"
    t.cell(1, 1).text = "Ô 2"
    doc.save(doc_path)

    importer = DocxImporter()
    res = importer.import_file(doc_path)
    assert len(res.document.blocks) == 1
    table_block = res.document.blocks[0]
    assert isinstance(table_block, Table)
    assert len(table_block.rows) == 2
    assert len(table_block.rows[0].cells) == 1
    assert table_block.rows[0].cells[0].colspan == 2
    assert len(table_block.rows[1].cells) == 2

