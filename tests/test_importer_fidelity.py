"""Kiểm thử tính trung thực và bảo toàn nội dung của các bộ nạp tài liệu (User Story 3 - L6, L7, L8, L10)."""
import pytest
from chuviettay.document.ir import Heading, LineBreak, ListBlock, Paragraph, Text
from chuviettay.importer.markdown_importer import MarkdownImporter

pytest.importorskip("docx", reason="Cần python-docx để chạy kiểm thử định dạng Word")


def test_markdown_preserves_inequalities():
    """Toán tử so sánh < và > không bị nuốt bởi bộ lọc HTML (L7)."""
    importer = MarkdownImporter()
    res = importer.import_text("So sánh: 1 < 2 và 3 > 2, ngoài ra $a < b$ và $c > d$")
    p = res.document.blocks[0]
    full_text = " ".join(getattr(inl, "text", "") or getattr(inl, "latex", "") for inl in p.inlines)
    assert "1 < 2" in full_text
    assert "3 > 2" in full_text
    assert "a < b" in full_text
    assert "c > d" in full_text

    # Kiểm tra trường hợp đặc biệt không có khoảng trắng quanh toán tử: $a<b$
    res2 = importer.import_text("$a<b$ và $c>d$, cũng như $a<b $ và $c>d$")
    p2 = res2.document.blocks[0]
    math_contents = [inl.latex for inl in p2.inlines if hasattr(inl, "latex")]
    assert "a<b" in math_contents
    assert "c>d" in math_contents


def test_markdown_nested_lists():
    """Danh sách Markdown lồng nhau nhiều cấp được duyệt đệ quy chính xác (L8)."""
    md_text = """
- Cấp 1 mục A
  - Cấp 2 mục A.1
  - Cấp 2 mục A.2
- Cấp 1 mục B
"""
    importer = MarkdownImporter()
    res = importer.import_text(md_text)
    assert len(res.document.blocks) == 1
    root_list = res.document.blocks[0]
    assert isinstance(root_list, ListBlock)
    assert len(root_list.items) == 2

    item_a_blocks = root_list.items[0]
    assert len(item_a_blocks) == 2
    assert isinstance(item_a_blocks[0], Paragraph)
    assert "Cấp 1 mục A" in item_a_blocks[0].inlines[0].text
    # Khối thứ 2 của item A là ListBlock lồng nhau
    assert isinstance(item_a_blocks[1], ListBlock)
    nested_list = item_a_blocks[1]
    assert len(nested_list.items) == 2


def test_markdown_code_block_and_hr_unsupported():
    """Khối mã (code block) và đường kẻ ngang (hr) không bị mất im lặng, được ghi vào unsupported (L8)."""
    md_text = """
Đoạn trước

```python
print("hello world")
```

---

Đoạn sau
"""
    importer = MarkdownImporter()
    res = importer.import_text(md_text)
    assert len(res.unsupported) >= 2
    assert any("code_block" in item for item in res.unsupported)
    assert any("hr" in item for item in res.unsupported)


def test_markdown_no_whitespace_before_punctuation():
    """Ranh giới inline kết thúc trước dấu câu không bị chèn khoảng trắng giả tạo (L6)."""
    importer = MarkdownImporter()
    res = importer.import_text("**chào**, các bạn!")
    p = res.document.blocks[0]
    assert len(p.inlines) == 1
    assert p.inlines[0].text == "chào, các bạn!"


def test_docx_runs_tabs_breaks_tracked_changes(tmp_path):
    """DOCX importer trích xuất toàn bộ text sau tab, xuống dòng mềm, ins, lọc delText và instrText (L10)."""
    import docx
    from docx.oxml import parse_xml
    from chuviettay.importer.docx_importer import DocxImporter

    doc_path = str(tmp_path / "comprehensive.docx")
    doc = docx.Document()

    # 1. Đoạn có tab, soft break và nhiều thẻ w:t trong cùng một run
    p1 = doc.add_paragraph()
    p1_xml = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '  <w:r>'
        '    <w:t>Trước tab</w:t>'
        '    <w:tab/>'
        '    <w:t>Sau tab</w:t>'
        '    <w:br/>'
        '    <w:t>Sau xuống dòng</w:t>'
        '  </w:r>'
        '</w:p>'
    )
    p1._element.getparent().replace(p1._element, parse_xml(p1_xml))

    # 2. Đoạn có tracked changes: w:ins (giữ) và w:del (bỏ qua)
    p2 = doc.add_paragraph()
    p2_xml = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '  <w:r><w:t>Văn bản gốc </w:t></w:r>'
        '  <w:ins><w:r><w:t>chữ mới chèn</w:t></w:r></w:ins>'
        '  <w:del><w:r><w:delText>chữ đã xoá</w:delText></w:r></w:del>'
        '</w:p>'
    )
    p2._element.getparent().replace(p2._element, parse_xml(p2_xml))

    # 3. Đoạn căn giữa và chứa field code w:instrText
    p3 = doc.add_paragraph()
    p3_xml = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '  <w:pPr><w:jc w:val="center"/></w:pPr>'
        '  <w:r><w:t>Trang số: </w:t></w:r>'
        '  <w:r><w:instrText>PAGE</w:instrText></w:r>'
        '</w:p>'
    )
    p3._element.getparent().replace(p3._element, parse_xml(p3_xml))

    doc.save(doc_path)

    importer = DocxImporter()
    res = importer.import_file(doc_path)

    assert len(res.document.blocks) == 3

    # Kiểm tra đoạn 1: có chữ trước tab, sau tab, và sau xuống dòng
    blk1 = res.document.blocks[0]
    assert any("Trước tab" in getattr(i, "text", "") for i in blk1.inlines)
    assert any("Sau tab" in getattr(i, "text", "") for i in blk1.inlines)
    assert any("Sau xuống dòng" in getattr(i, "text", "") for i in blk1.inlines)
    assert any(isinstance(i, LineBreak) for i in blk1.inlines)

    # Kiểm tra đoạn 2: có "chữ mới chèn", KHÔNG có "chữ đã xoá"
    blk2 = res.document.blocks[1]
    blk2_text = "".join(getattr(i, "text", "") for i in blk2.inlines)
    assert "chữ mới chèn" in blk2_text
    assert "chữ đã xoá" not in blk2_text

    # Kiểm tra đoạn 3: căn giữa, có "Trang số: ", instrText ghi vào unsupported
    blk3 = res.document.blocks[2]
    assert blk3.align == "center"
    blk3_text = "".join(getattr(i, "text", "") for i in blk3.inlines)
    assert "Trang số: " in blk3_text
    assert "PAGE" not in blk3_text
    assert any("field_instruction" in u for u in res.unsupported)
