"""Fidelity: bản làm trắng phải làm trắng cả công thức OMML, nếu không công thức in đen nằm dưới chữ viết tay."""
import zipfile

import docx
from docx.oxml import parse_xml
from lxml import etree

from chuviettay.fidelity.background import M_NS, W_NS, WhiteoutBackgroundGenerator, _whiten_xml_bytes

NS = f'xmlns:m="{M_NS}" xmlns:w="{W_NS}"'


def q(ns, tag):
    return f"{{{ns}}}{tag}"


def build_docx(tmp_path, inner, name="eq.docx"):
    d = docx.Document()
    d.add_paragraph("Chữ thường")
    p = d.add_paragraph("Công thức: ")
    p._element.append(parse_xml(f"<m:oMath {NS}>{inner}</m:oMath>"))
    path = str(tmp_path / name)
    d.save(path)
    return path


def whiten(tmp_path, inner):
    src = build_docx(tmp_path, inner)
    out = str(tmp_path / "white.docx")
    WhiteoutBackgroundGenerator.create_whiteout_docx(src, out)
    return etree.fromstring(zipfile.ZipFile(out).read("word/document.xml"))


def color_of(wrpr):
    c = wrpr.find(q(W_NS, "color"))
    return None if c is None else c.get(q(W_NS, "val"))


def test_math_runs_are_whitened(tmp_path):
    root = whiten(tmp_path, "<m:r><m:t>x+1</m:t></m:r>")
    runs = root.findall(f".//{q(M_NS, 'r')}")
    assert runs and all(color_of(r.find(q(W_NS, "rPr"))) == "FFFFFF" for r in runs)


def test_math_run_child_order_follows_schema(tmp_path):
    """m:r phải có thứ tự m:rPr, w:rPr, m:t."""
    root = whiten(tmp_path, '<m:r><m:rPr><m:sty m:val="p"/></m:rPr><m:t>sin</m:t></m:r><m:r><m:t>x</m:t></m:r>')
    for r in root.findall(f".//{q(M_NS, 'r')}"):
        names = [c.tag for c in r]
        assert names == ([q(M_NS, "rPr"), q(W_NS, "rPr"), q(M_NS, "t")] if len(names) == 3 else [q(W_NS, "rPr"), q(M_NS, "t")])


def test_existing_math_run_formatting_is_kept_and_recoloured(tmp_path):
    inner = '<m:r><w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/><w:color w:val="FF0000" w:themeColor="accent1"/><w:sz w:val="28"/></w:rPr><m:t>x</m:t></m:r>'
    root = whiten(tmp_path, inner)
    rpr = root.find(f".//{q(M_NS, 'r')}/{q(W_NS, 'rPr')}")
    assert rpr.find(q(W_NS, "rFonts")) is not None and rpr.find(q(W_NS, "sz")) is not None
    color = rpr.find(q(W_NS, "color"))
    assert color.get(q(W_NS, "val")) == "FFFFFF" and color.get(q(W_NS, "themeColor")) is None
    children = [c.tag.split("}")[-1] for c in rpr]
    assert children.index("color") < children.index("sz")        # w:color phải đứng trước w:sz theo schema


def test_structure_control_properties_are_whitened(tmp_path):
    """Gạch phân số, dấu căn, ngoặc... lấy màu từ m:ctrlPr."""
    inner = ("<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>"
             "<m:d><m:e><m:r><m:t>x</m:t></m:r></m:e></m:d>")
    root = whiten(tmp_path, inner)
    ctrls = root.findall(f".//{q(M_NS, 'ctrlPr')}")
    assert len(ctrls) == 2 and all(color_of(c.find(q(W_NS, "rPr"))) == "FFFFFF" for c in ctrls)
    for pr in root.findall(f".//{q(M_NS, 'fPr')}") + root.findall(f".//{q(M_NS, 'dPr')}"):
        assert pr[-1].tag == q(M_NS, "ctrlPr")                    # ctrlPr luôn là phần tử cuối của *Pr


def test_ctrlpr_is_added_after_existing_properties(tmp_path):
    inner = '<m:d><m:dPr><m:begChr m:val="["/><m:endChr m:val="]"/></m:dPr><m:e><m:r><m:t>x</m:t></m:r></m:e></m:d>'
    root = whiten(tmp_path, inner)
    pr = root.find(f".//{q(M_NS, 'dPr')}")
    assert [c.tag.split("}")[-1] for c in pr] == ["begChr", "endChr", "ctrlPr"]


def test_plain_text_runs_still_whitened_and_graphics_untouched():
    xml = (f'<w:document {NS}><w:body><w:p><w:r><w:t>abc</w:t></w:r></w:p></w:body></w:document>').encode()
    root = etree.fromstring(_whiten_xml_bytes(xml))
    assert color_of(root.find(f".//{q(W_NS, 'r')}/{q(W_NS, 'rPr')}")) == "FFFFFF"


def test_whitened_package_still_opens_and_text_is_intact(tmp_path):
    src = build_docx(tmp_path, "<m:f><m:num><m:r><m:t>a+1</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>")
    out = str(tmp_path / "w.docx")
    WhiteoutBackgroundGenerator.create_whiteout_docx(src, out)
    d = docx.Document(out)
    assert [p.text for p in d.paragraphs][:2] == ["Chữ thường", "Công thức: "]
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    assert "a+1" in xml                                           # nội dung công thức còn nguyên, chỉ đổi màu


def test_idempotent_whitening(tmp_path):
    src = build_docx(tmp_path, "<m:r><m:t>x</m:t></m:r>")
    once, twice = str(tmp_path / "1.docx"), str(tmp_path / "2.docx")
    WhiteoutBackgroundGenerator.create_whiteout_docx(src, once)
    WhiteoutBackgroundGenerator.create_whiteout_docx(once, twice)
    a = etree.fromstring(zipfile.ZipFile(once).read("word/document.xml"))
    b = etree.fromstring(zipfile.ZipFile(twice).read("word/document.xml"))
    assert etree.tostring(a, method="c14n") == etree.tostring(b, method="c14n")
