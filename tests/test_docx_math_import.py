"""Nạp .docx có công thức: OMML (inline/display/bảng), đối tượng MathType (OLE) và trường EQ cũ."""
import docx
from docx.oxml import parse_xml

from chuviettay.document.ir import MathBlock, MathInline, Paragraph, Table
from chuviettay.importer.docx_importer import DocxImporter
from chuviettay.importer.omml import M_NS
from chuviettay.layout.math_layout import MathLayoutEngine
from chuviettay.math import MathRow, NAry, SymbolNode, parse_latex_math, walk
from tests.mathtest_helpers import canon

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = f'xmlns:m="{M_NS}" xmlns:w="{W}" xmlns:v="urn:schemas-microsoft-com:vml" ' \
     f'xmlns:o="urn:schemas-microsoft-com:office:office" ' \
     f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'


def mr(t):
    return f"<m:r><m:t>{t}</m:t></m:r>"


def omath(inner):
    return parse_xml(f"<m:oMath {NS}>{inner}</m:oMath>")


def opara(*inners):
    body = "".join(f"<m:oMath>{i}</m:oMath>" for i in inners)
    return parse_xml(f"<m:oMathPara {NS}>{body}</m:oMathPara>")


def ole_run(progid="Equation.DSMT4"):
    return parse_xml(
        f'<w:r {NS}><w:object w:dxaOrig="1440" w:dyaOrig="440">'
        '<v:shape id="_x0000_i1025" type="#_x0000_t75" style="width:72pt;height:22pt"><v:imagedata r:id="rId5" o:title=""/></v:shape>'
        f'<o:OLEObject Type="Embed" ProgID="{progid}" ShapeID="_x0000_i1025" DrawAspect="Content" ObjectID="_1" r:id="rId6"/>'
        "</w:object></w:r>")


def run_import(tmp_path, build):
    d = docx.Document()
    build(d)
    path = str(tmp_path / "t.docx")
    d.save(path)
    return DocxImporter().import_file(path)


def math_inlines(block):
    return [i for i in getattr(block, "inlines", []) if isinstance(i, MathInline)]


def test_inline_equation_inside_text(tmp_path):
    def build(d):
        p = d.add_paragraph("Giải phương trình ")
        p._element.append(omath(mr("2x+3=0")))
        p.add_run(" ta được x = -3/2.")

    res = run_import(tmp_path, build)
    para = res.document.blocks[0]
    assert isinstance(para, Paragraph)
    kinds = [type(i).__name__ for i in para.inlines]
    assert kinds == ["Text", "MathInline", "Text"]
    math = para.inlines[1]
    assert math.latex and canon(math.ast) == canon(parse_latex_math("2x+3=0"))
    assert res.unsupported == [] and res.warnings == []


def test_paragraph_with_only_an_equation_is_a_math_block(tmp_path):
    res = run_import(tmp_path, lambda d: d.add_paragraph()._element.append(omath(mr("x^2"))))
    assert len(res.document.blocks) == 1 and isinstance(res.document.blocks[0], MathBlock)


def test_display_equation_between_text_is_split_into_blocks(tmp_path):
    def build(d):
        p = d.add_paragraph("Trước đó ")
        p._element.append(opara(mr("a=b"), mr("c=d")))
        p.add_run(" và sau đó.")

    res = run_import(tmp_path, build)
    types = [type(b).__name__ for b in res.document.blocks]
    assert types == ["Paragraph", "MathBlock", "MathBlock", "Paragraph"]
    assert res.document.blocks[0].inlines[0].text.startswith("Trước đó")
    assert res.document.blocks[3].inlines[0].text.strip() == "và sau đó."


def test_equation_in_table_cell_is_kept(tmp_path):
    def build(d):
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).paragraphs[0].add_run("Công thức")
        t.cell(0, 1).paragraphs[0]._element.append(omath(f"<m:f><m:num>{mr('a')}</m:num><m:den>{mr('b')}</m:den></m:f>"))

    res = run_import(tmp_path, build)
    table = next(b for b in res.document.blocks if isinstance(b, Table))
    blocks = table.rows[0].cells[1].blocks
    math = [b for b in blocks if isinstance(b, MathBlock)] + [i for b in blocks for i in getattr(b, "inlines", []) if isinstance(i, MathInline)]
    assert math, "công thức trong ô bảng bị mất"


def test_empty_equation_is_ignored(tmp_path):
    def build(d):
        p = d.add_paragraph("Chữ ")
        p._element.append(omath(""))

    res = run_import(tmp_path, build)
    assert not any(isinstance(i, MathInline) for b in res.document.blocks for i in getattr(b, "inlines", []))


def test_nary_and_matrix_are_real_nodes(tmp_path):
    def build(d):
        p = d.add_paragraph("x ")
        p._element.append(omath(
            f"<m:nary><m:naryPr><m:chr m:val=\"∑\"/></m:naryPr><m:sub>{mr('i=1')}</m:sub><m:sup>{mr('n')}</m:sup><m:e>{mr('i')}</m:e></m:nary>"))

    res = run_import(tmp_path, build)
    math = math_inlines(res.document.blocks[0])[0]
    assert any(isinstance(n, NAry) and n.op == "∑" for n in walk(math.ast))
    assert res.unsupported == []


# ------------------------------------------------------------------ MathType (OLE) và trường EQ
def test_mathtype_object_is_never_silently_dropped(tmp_path):
    """Lỗi cũ: w:object bị bỏ -> 'Giải phương trình  với x thuộc R.' và không có cảnh báo nào."""
    def build(d):
        p = d.add_paragraph("Giải phương trình ")
        p._element.append(ole_run())
        p.add_run(" với x thuộc R.")

    res = run_import(tmp_path, build)
    para = res.document.blocks[0]
    assert [type(i).__name__ for i in para.inlines] == ["Text", "MathInline", "Text"]
    assert para.inlines[1].ast == MathRow([SymbolNode("□")])
    assert len(res.warnings) == 1 and "1 công thức MathType" in res.warnings[0] and "Convert Equations" in res.warnings[0]
    assert any(u.startswith("mathtype:") and "Equation.DSMT4" in u for u in res.unsupported)


def test_paragraph_made_only_of_mathtype_is_kept(tmp_path):
    def build(d):
        d.add_paragraph()._element.append(ole_run("Equation.3"))
        d.add_paragraph()._element.append(ole_run("MathType.Equation"))

    res = run_import(tmp_path, build)
    assert len(res.document.blocks) == 2
    assert "2 công thức MathType" in res.warnings[0]


def test_mathtype_count_resets_between_imports(tmp_path):
    imp = DocxImporter()
    d = docx.Document()
    d.add_paragraph()._element.append(ole_run())
    path = str(tmp_path / "a.docx")
    d.save(path)
    first = imp.import_file(path)
    second = imp.import_file(path)
    assert "1 công thức" in first.warnings[0] and "1 công thức" in second.warnings[0]


def test_mathtype_inside_table_cell_is_counted(tmp_path):
    def build(d):
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).paragraphs[0]._element.append(ole_run())

    res = run_import(tmp_path, build)
    assert res.warnings and "1 công thức MathType" in res.warnings[0]


def test_other_ole_objects_are_reported_but_not_as_mathtype(tmp_path):
    def build(d):
        p = d.add_paragraph("Bảng ")
        p._element.append(ole_run("Excel.Sheet.12"))

    res = run_import(tmp_path, build)
    assert res.warnings == []
    assert any("Excel.Sheet.12" in u for u in res.unsupported)
    assert not any(isinstance(i, MathInline) for b in res.document.blocks for i in getattr(b, "inlines", []))


def test_legacy_eq_field_is_reported(tmp_path):
    def build(d):
        p = d.add_paragraph("Phân số ")
        p._element.append(parse_xml(f'<w:r {NS}><w:fldChar w:fldCharType="begin"/></w:r>'))
        p._element.append(parse_xml(f'<w:r {NS}><w:instrText xml:space="preserve"> EQ \\f(1,2) </w:instrText></w:r>'))
        p._element.append(parse_xml(f'<w:r {NS}><w:fldChar w:fldCharType="end"/></w:r>'))

    res = run_import(tmp_path, build)
    assert res.warnings and "trường EQ" in res.warnings[0]
    assert any("EQ" in u for u in res.unsupported)


def test_mathtype_placeholder_renders_as_box_and_is_reported_missing(tmp_path):
    from tests.mathtest_helpers import make_block_bank
    from chuviettay.model.bank import Bank
    import random

    bank = Bank(make_block_bank(str(tmp_path / "b.json.gz")))
    from chuviettay.model.writer import Writer

    eng = MathLayoutEngine(bank, writer=Writer(bank, random.Random(1)), rnd=random.Random(1))
    item = eng.measure(MathRow([SymbolNode("□")]))
    assert item.strokes and item.size.width > 0 and eng.missing_symbols == {"□": 1}
