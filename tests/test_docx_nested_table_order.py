"""Kiểm thử thứ tự xuất hiện tài liệu (document order) khi ô bảng chứa xen kẽ
đoạn văn và bảng lồng nhau (nested table).
"""
import docx
import pytest

from chuviettay.document.ir import Paragraph, Table
from chuviettay.importer.docx_importer import DocxImporter


def test_docx_nested_table_preserves_document_block_sequence(tmp_path):
    """Kiểm tra: Paragraph A -> Nested Table -> Paragraph B giữ đúng thứ tự xuất hiện
    thay vì bị gom tất cả paragraphs trước rồi mới nối nested table ở cuối.
    """
    docx_doc = docx.Document()
    main_table = docx_doc.add_table(rows=1, cols=1)
    cell = main_table.cell(0, 0)

    # 1. Paragraph đầu tiên
    cell.paragraphs[0].text = "Đoạn văn A"

    # 2. Bảng lồng nhau
    nested = cell.add_table(rows=1, cols=2)
    nested.cell(0, 0).paragraphs[0].text = "Con 1"
    nested.cell(0, 1).paragraphs[0].text = "Con 2"

    # 3. Paragraph thứ hai
    cell.add_paragraph("Đoạn văn B")

    docx_file = str(tmp_path / "nested_order.docx")
    docx_doc.save(docx_file)

    importer = DocxImporter()
    res = importer.import_file(docx_file)

    assert len(res.document.blocks) >= 1
    table_block = res.document.blocks[0]
    assert isinstance(table_block, Table)

    cell_ir = table_block.rows[0].cells[0]
    assert len(cell_ir.blocks) == 3, f"Ô phải có đúng 3 block (A -> Nested Table -> B), nhận được {len(cell_ir.blocks)}"

    # Trích xuất văn bản của từng block
    def block_text(b):
        if hasattr(b, "inlines"):
            return " ".join(inl.text for inl in b.inlines if hasattr(inl, "text") and inl.text)
        return ""

    texts = [block_text(b) for b in cell_ir.blocks]
    assert "Đoạn văn A" in texts[0], f"Block 0 phải là Đoạn văn A, nhận được: {texts[0]}"
    assert "Con 1" in texts[1] and "Con 2" in texts[1], f"Block 1 phải là bảng con, nhận được: {texts[1]}"
    assert "Đoạn văn B" in texts[2], f"Block 2 phải là Đoạn văn B, nhận được: {texts[2]}"
