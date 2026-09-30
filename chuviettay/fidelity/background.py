"""Bộ tạo tệp nền tài liệu không chứa chữ in (Whiteout Background Generator).

Thực hiện biến đổi Whiteout Run Transform: đổi toàn bộ chữ in thành màu trắng (#FFFFFF),
giữ nguyên 100% hình ảnh (inline drawing), đường viền bảng, biểu đồ và tỷ lệ bố cục.
"""
from __future__ import annotations

import logging
import os
import tempfile

_log = logging.getLogger(__name__)


class WhiteoutBackgroundGenerator:
    """Tạo bản sao DOCX đã làm trắng chữ in để làm nền cho Xournal++."""

    @classmethod
    def create_whiteout_docx(cls, docx_path: str, output_path: str | None = None) -> str:
        """Đổi toàn bộ text runs trong DOCX thành màu trắng và lưu ra file mới."""
        docx_abs = os.path.abspath(docx_path)
        if not os.path.exists(docx_abs):
            raise FileNotFoundError(f"Không tìm thấy tệp DOCX: {docx_path}")

        if output_path is None:
            fd, output_path = tempfile.mkstemp(prefix="whiteout_", suffix=".docx")
            os.close(fd)

        try:
            import docx
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            from docx.shared import RGBColor

            doc = docx.Document(docx_abs)
            white = RGBColor(255, 255, 255)

            # 1. Làm trắng toàn bộ runs thông qua API python-docx cấp cao (paragraphs & tables)
            for p in doc.paragraphs:
                for r in p.runs:
                    r.font.color.rgb = white

            def _whiten_table_cells(table) -> None:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.font.color.rgb = white
                        # Xử lý đệ quy bảng lồng nhau bên trong ô
                        for sub_table in getattr(cell, "tables", []):
                            _whiten_table_cells(sub_table)

            for tbl in doc.tables:
                _whiten_table_cells(tbl)

            # 2. Làm trắng toàn diện cấp OpenXML DOM (quét mọi <w:r> kể cả trong DrawingML textboxes & shapes)
            for r in doc.element.xpath(".//w:r"):
                try:
                    rPr = r.get_or_add_rPr()
                    color = rPr.find(qn("w:color"))
                    if color is None:
                        color = OxmlElement("w:color")
                        rPr.append(color)
                    color.set(qn("w:val"), "FFFFFF")
                except Exception:
                    pass

            # 3. Làm trắng DrawingML text runs (<a:r> trong wp:inline / txBody)
            for a_r in doc.element.xpath(".//a:r"):
                try:
                    a_rPr = a_r.find(qn("a:rPr"))
                    if a_rPr is None:
                        a_rPr = OxmlElement("a:rPr")
                        a_r.insert(0, a_rPr)
                    for child in list(a_rPr):
                        if child.tag.endswith(("solidFill", "gradFill", "blipFill", "pattFill", "noFill")):
                            a_rPr.remove(child)
                    solid_fill = OxmlElement("a:solidFill")
                    srgb_clr = OxmlElement("a:srgbClr")
                    srgb_clr.set("val", "FFFFFF")
                    solid_fill.append(srgb_clr)
                    a_rPr.append(solid_fill)
                except Exception:
                    pass

            # 4. Làm trắng tất cả headers và footers (kể cả first page, even page) cấp DOM
            for sec in doc.sections:
                for hf in (
                    getattr(sec, "header", None),
                    getattr(sec, "footer", None),
                    getattr(sec, "first_page_header", None),
                    getattr(sec, "first_page_footer", None),
                    getattr(sec, "even_page_header", None),
                    getattr(sec, "even_page_footer", None),
                ):
                    if hf is not None:
                        for p in hf.paragraphs:
                            for r in p.runs:
                                r.font.color.rgb = white
                        if hasattr(hf, "part") and hf.part is not None:
                            try:
                                for r in hf.part.element.xpath(".//w:r"):
                                    rPr = r.get_or_add_rPr()
                                    color = rPr.find(qn("w:color"))
                                    if color is None:
                                        color = OxmlElement("w:color")
                                        rPr.append(color)
                                    color.set(qn("w:val"), "FFFFFF")
                            except Exception:
                                pass

            # 5. Quét tất cả các related parts của document (footnotes, endnotes, comments...)
            if hasattr(doc, "part") and hasattr(doc.part, "related_parts"):
                for rel_part in doc.part.related_parts.values():
                    if hasattr(rel_part, "element"):
                        try:
                            for r in rel_part.element.xpath(".//w:r"):
                                rPr = r.get_or_add_rPr()
                                color = rPr.find(qn("w:color"))
                                if color is None:
                                    color = OxmlElement("w:color")
                                    rPr.append(color)
                                color.set(qn("w:val"), "FFFFFF")
                        except Exception:
                            pass

            doc.save(output_path)
            return output_path

        except Exception as e:
            _log.error("Lỗi khi tạo whiteout DOCX: %s", e)
            if output_path and os.path.exists(output_path):
                try:
                    os.remove(output_path)
                except OSError:
                    pass
            raise RuntimeError(f"Không thể làm trắng chữ in trong DOCX ({docx_path}): {e}") from e
